"""Preview and import order spreadsheets without changing task progress."""

import csv
import io
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime, timezone as datetime_timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zoneinfo import ZoneInfo

from django.conf import settings
from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_datetime
from openpyxl.utils.exceptions import InvalidFileException

from .models import ImportBatch, Order


MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
HEADERS = {
    "order id": "order_id",
    "order date": "order_date",
    "customer name": "customer_name",
    "phone number": "phone",
    "phone": "phone",
    "order value": "order_value",
    "payment type": "payment_type",
    "order status": "status",
    "status": "status",
    "dispatch date": "dispatch_date",
    "expected delivery date": "expected_delivery_date",
    "assigned employee": "assigned_employee",
    "confirmed at": "confirmed_at",
    "confirmation date": "confirmed_at",
    "confirmed date": "confirmed_at",
}
REQUIRED_HEADERS = {
    "order_id", "order_date", "customer_name", "phone", "order_value",
    "payment_type", "status", "dispatch_date", "expected_delivery_date",
    "assigned_employee",
}
STATUS_ALIASES = {
    "unconfirmed": "unconfirmed", "pending": "unconfirmed",
    "pending confirmation": "unconfirmed", "confirmed": "confirmed",
    "dispatched": "dispatched", "shipped": "dispatched",
    "in transit": "dispatched", "delivery failed": "delivery_failed",
    "failed delivery": "delivery_failed", "failed delivery attempt": "delivery_failed",
    "delivered": "delivered", "cancelled": "cancelled", "canceled": "cancelled",
}
# These are the source fields used to detect an identical re-upload.
ORDER_FIELDS = (
    "order_id", "order_date", "ordered_at", "customer_name", "phone",
    "order_value", "payment_type", "status", "confirmed_at", "dispatch_date",
    "dispatch_at", "expected_delivery_date", "expected_delivery_at",
    "assigned_employee_id", "data_issues",
)
DATE_FIELDS = {"order_date", "dispatch_date", "expected_delivery_date"}
TIME_FIELDS = {"ordered_at", "confirmed_at", "dispatch_at", "expected_delivery_at"}


def clean_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def words(value):
    return " ".join(clean_text(value).lower().replace("_", " ").replace("-", " ").split())


def parse_date(value, label):
    """Return (calendar date, precise timestamp), preserving date-only input."""
    if value is None or clean_text(value) == "":
        return None, None
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        return value.isoformat(), None
    else:
        text = clean_text(value)
        for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(text, pattern).date().isoformat(), None
            except ValueError:
                pass
        try:
            if not re.match(r"^\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}", text):
                raise ValueError
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            parsed = None
            for pattern in ("%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S"):
                try:
                    parsed = datetime.strptime(text, pattern)
                    break
                except ValueError:
                    pass
            if parsed is None:
                raise ValueError(f"{label}: invalid date. Use YYYY-MM-DD or a full timestamp.")
    local_zone = ZoneInfo(settings.TIME_ZONE)
    try:
        if timezone.is_naive(parsed):
            parsed = timezone.make_aware(parsed, local_zone)
        calendar_date = parsed.astimezone(local_zone).date().isoformat()
        return calendar_date, parsed.astimezone(datetime_timezone.utc).isoformat()
    except (OverflowError, ValueError) as exc:
        raise ValueError(f"{label}: timestamp is outside the supported date range.") from exc


def snapshot(order):
    data = {}
    for field in ORDER_FIELDS:
        value = getattr(order, field)
        if isinstance(value, datetime):
            value = value.astimezone(datetime_timezone.utc).isoformat()
        elif isinstance(value, date):
            value = value.isoformat()
        elif isinstance(value, Decimal):
            value = format(value, ".2f")
        data[field] = value
    return data


def _read_file(uploaded_file):
    name = Path(uploaded_file.name).name
    suffix = Path(name).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ValueError("Please upload a CSV or .xlsx file. Older .xls files are not supported.")
    uploaded_file.seek(0)
    content = uploaded_file.read(MAX_BYTES + 1)
    if len(content) > MAX_BYTES:
        raise ValueError("The file is too large. Maximum upload size is 5 MB.")
    rows = []
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text, newline=""), dialect)
            header = next(reader, [])
            for row in reader:
                if any(clean_text(value) for value in row):
                    rows.append((reader.line_num, row, []))
                    if len(rows) > MAX_ROWS:
                        raise ValueError("A file may contain at most 10,000 non-empty order rows.")
        else:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("The expanded workbook exceeds the 50 MB safety limit.")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise ValueError("A workbook may contain at most 10,000 order rows.")
                if sheet.max_column and sheet.max_column > 100:
                    raise ValueError("A workbook may contain at most 100 columns.")
                # Some spreadsheet writers save incorrect worksheet dimensions.
                sheet.reset_dimensions()
                iterator = sheet.iter_rows()
                header_cells = next(iterator, [])
                if any(cell.data_type == "f" for cell in header_cells):
                    raise ValueError("Formula cells are not allowed in column headers.")
                header = [cell.value for cell in header_cells]
                for number, cells in enumerate(iterator, 2):
                    if len(cells) > 100:
                        raise ValueError("A workbook may contain at most 100 columns.")
                    values = [
                        cell.value.date() if isinstance(cell.value, datetime)
                        and is_datetime(cell.number_format) == "date" else cell.value
                        for cell in cells
                    ]
                    if any(clean_text(value) for value in values):
                        errors = ["Formula cells are not allowed; upload plain values instead."] if any(
                            cell.data_type == "f" for cell in cells
                        ) else []
                        rows.append((number, values, errors))
                        if len(rows) > MAX_ROWS:
                            raise ValueError("A workbook may contain at most 10,000 non-empty order rows.")
            finally:
                workbook.close()
    except (UnicodeDecodeError, csv.Error, zipfile.BadZipFile, KeyError, OSError, ParseError, InvalidFileException, IndexError) as exc:
        raise ValueError("The file could not be read. Use a UTF-8 CSV or a valid .xlsx workbook.") from exc
    if not header or not rows:
        raise ValueError("The file must include a header row and at least one order.")
    if len(header) > 100:
        raise ValueError("A file may contain at most 100 columns.")
    canonical = [HEADERS.get(words(value)) for value in header]
    recognized = [field for field in canonical if field]
    if len(recognized) != len(set(recognized)):
        raise ValueError("The file has duplicate column headers.")
    missing = sorted(REQUIRED_HEADERS - set(recognized))
    if missing:
        raise ValueError("Missing columns: " + ", ".join(field.replace("_", " ") for field in missing) + ".")
    return name, canonical, rows


def _normalize(row, user, employees):
    errors, warnings = [], []
    data = {"order_id": clean_text(row.get("order_id"))}
    for field, label, maximum in (
        ("order_id", "Order ID", 64), ("customer_name", "Customer name", 100),
        ("phone", "Phone number", 32),
    ):
        value = clean_text(row.get(field))
        data[field] = value
        if not value:
            errors.append(f"{label} is required.")
        elif len(value) > maximum:
            errors.append(f"{label} is too long (maximum {maximum} characters).")
    phone = data.get("phone", "")
    if phone and (not re.fullmatch(r"[+\d() .-]+", phone) or not 7 <= sum(c.isdigit() for c in phone) <= 15):
        errors.append("Phone number must contain 7–15 digits; +, spaces and separators are allowed.")

    try:
        amount = Decimal(clean_text(row.get("order_value")).replace("₹", "").replace(",", ""))
        if not amount.is_finite() or amount < 0 or amount > Decimal("9999999999.99"):
            raise InvalidOperation
        if amount * 100 != (amount * 100).to_integral_value():
            raise InvalidOperation
        data["order_value"] = format(amount, ".2f")
    except (InvalidOperation, ValueError):
        errors.append("Order value must be a non-negative amount with at most two decimal places.")

    payment = words(row.get("payment_type"))
    payment = {"prepaid": "prepaid", "cod": "cod", "cash on delivery": "cod"}.get(payment)
    if payment is None:
        errors.append("Payment type must be prepaid or COD.")
    data["payment_type"] = payment
    status = STATUS_ALIASES.get(words(row.get("status")))
    if status is None:
        errors.append("Order status is missing or unknown. Use unconfirmed, confirmed, dispatched, delivery_failed, delivered or cancelled.")
    data["status"] = status

    parsed_dates = {}
    for source, target_date, target_time, label in (
        ("order_date", "order_date", "ordered_at", "Order date"),
        ("confirmed_at", None, "confirmed_at", "Confirmed at"),
        ("dispatch_date", "dispatch_date", "dispatch_at", "Dispatch date"),
        ("expected_delivery_date", "expected_delivery_date", "expected_delivery_at", "Expected delivery date"),
    ):
        try:
            day, moment = parse_date(row.get(source), label)
            parsed_dates[source] = (day, moment)
            if source == "order_date" and day is None:
                errors.append("Order date is required.")
            if source == "confirmed_at" and day and not moment:
                warnings.append("Confirmation time is missing; provide a full Confirmed at timestamp for the dispatch rule.")
            if target_date:
                data[target_date] = day
            data[target_time] = moment
        except (ValueError, OverflowError) as exc:
            errors.append(str(exc))

    order_day, order_time = parsed_dates.get("order_date", (None, None))
    confirm_day, confirm_time = parsed_dates.get("confirmed_at", (None, None))
    dispatch_day, dispatch_time = parsed_dates.get("dispatch_date", (None, None))
    delivery_day, delivery_time = parsed_dates.get("expected_delivery_date", (None, None))
    for label, day, moment in (
        ("Confirmation", confirm_day, confirm_time), ("Dispatch", dispatch_day, dispatch_time),
        ("Expected delivery", delivery_day, delivery_time),
    ):
        if order_day and day and (day < order_day or (moment and order_time and moment < order_time)):
            errors.append(f"{label} cannot be before the order date.")
    if confirm_day and dispatch_day and (
        dispatch_day < confirm_day or (dispatch_time and confirm_time and dispatch_time < confirm_time)
    ):
        errors.append("Dispatch cannot be before confirmation.")
    if dispatch_day and delivery_day and (
        delivery_day < dispatch_day or (delivery_time and dispatch_time and delivery_time < dispatch_time)
    ):
        errors.append("Expected delivery cannot be before dispatch.")
    if status in {"unconfirmed", "confirmed"} and dispatch_day:
        errors.append("An unconfirmed or confirmed order cannot already have a dispatch date; correct its status.")

    employee_name = clean_text(row.get("assigned_employee"))
    matching = employees.get(employee_name.casefold(), []) if employee_name else []
    if len(matching) == 1:
        data["assigned_employee_id"] = matching[0]
    else:
        data["assigned_employee_id"] = user.pk
        warnings.append("Assigned employee is missing, unknown or ambiguous; assigned to the importing manager until corrected.")
    data["data_issues"] = warnings.copy()
    return data, errors, warnings


def preview_import(uploaded_file, user):
    """Save a validation preview only; no orders or tasks are changed here."""
    if not user.is_staff:
        raise ValueError("Only a manager can import orders.")
    name, headers, rows = _read_file(uploaded_file)
    employees = defaultdict(list)
    for employee in get_user_model().objects.filter(is_active=True):
        for label in {employee.username.casefold(), employee.get_full_name().strip().casefold()} - {""}:
            employees[label].append(employee.pk)
    preview, issues = [], []
    for number, values, file_errors in rows:
        row = {key: values[index] if index < len(values) else None for index, key in enumerate(headers) if key}
        data, errors, warnings = _normalize(row, user, employees)
        errors = file_errors + errors
        if len(values) > len(headers):
            errors.append("This row has more values than the header row.")
        preview.append({
            "row_number": number, "action": "rejected" if errors else "create",
            "data": data, "errors": errors, "warnings": warnings, "before": None,
        })

    groups = defaultdict(list)
    for item in preview:
        if item["data"].get("order_id"):
            groups[item["data"]["order_id"]].append(item)
    for items in groups.values():
        if len(items) < 2:
            continue
        identical = all(item["data"] == items[0]["data"] and not item["errors"] for item in items)
        if identical:
            for item in items[1:]:
                item["action"] = "duplicate"
                item["warnings"].append("Identical duplicate order ID in this file; only the first row will be used.")
        else:
            for item in items:
                item["action"] = "rejected"
                item["errors"].append("Conflicting duplicate order ID in this file; correct every row with this ID.")

    existing = {order.order_id: order for order in Order.objects.filter(order_id__in=groups)}
    for item in preview:
        if item["action"] == "create" and item["data"]["order_id"] in existing:
            item["before"] = snapshot(existing[item["data"]["order_id"]])
            item["action"] = "unchanged" if item["before"] == item["data"] else "update"
            item["warnings"].append(
                "Existing order ID: identical data, so it will be left unchanged." if item["action"] == "unchanged"
                else "Existing order ID: changed data needs explicit update confirmation."
            )
        for severity, messages in (("error", item["errors"]), ("warning", item["warnings"])):
            for message in messages:
                issues.append({"row_number": item["row_number"], "order_id": item["data"].get("order_id", ""),
                               "field": "row", "message": message, "severity": severity})
    summary = {
        "rows": len(preview), "valid": sum(item["action"] in {"create", "update", "unchanged"} for item in preview),
        "rejected": sum(item["action"] == "rejected" for item in preview),
        "duplicates": sum(item["action"] == "duplicate" for item in preview),
        "created": sum(item["action"] == "create" for item in preview),
        "updated": sum(item["action"] == "update" for item in preview),
        "unchanged": sum(item["action"] == "unchanged" for item in preview), "skipped_updates": 0,
    }
    return ImportBatch.objects.create(filename=name, uploaded_by=user, status="preview", preview=preview, issues=issues, summary=summary)


def _model_values(data):
    values = data.copy()
    for field in DATE_FIELDS:
        values[field] = date.fromisoformat(data[field]) if data.get(field) else None
    for field in TIME_FIELDS:
        values[field] = datetime.fromisoformat(data[field]) if data.get(field) else None
    values["order_value"] = Decimal(data["order_value"])
    return values


@transaction.atomic
def commit_import(batch, user, allow_updates=False):
    """Apply an owned preview once, checking that its existing-order snapshot is current."""
    batch = ImportBatch.objects.select_for_update().get(pk=batch.pk)
    if not user.is_staff or batch.uploaded_by_id != user.pk:
        raise ValueError("Only the manager who uploaded this file can confirm it.")
    if batch.status == "completed":
        return batch.summary
    if batch.status != "preview":
        raise ValueError("This import is not awaiting confirmation.")
    summary = batch.summary.copy()
    summary.update(created=0, updated=0, unchanged=0, skipped_updates=0)
    for item in batch.preview:
        action = item["action"]
        if action in {"rejected", "duplicate"}:
            continue
        data = item["data"]
        existing = Order.objects.filter(order_id=data["order_id"]).first()
        if action == "create":
            if existing is not None:
                raise ValueError("An order changed after this preview. Upload the file again to review current data.")
            Order.objects.create(**_model_values(data))
            summary["created"] += 1
        else:
            if existing is None or snapshot(existing) != item["before"]:
                raise ValueError("An order changed after this preview. Upload the file again to review current data.")
            if action == "unchanged":
                summary["unchanged"] += 1
            elif not allow_updates:
                summary["skipped_updates"] += 1
            else:
                for field, value in _model_values(data).items():
                    setattr(existing, field, value)
                existing.save()
                summary["updated"] += 1
    batch.status = "completed"
    batch.summary = summary
    batch.save(update_fields=["status", "summary"])
    return summary
